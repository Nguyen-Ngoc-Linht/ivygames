import openpyxl
import json
import re

wb = openpyxl.load_workbook(r'D:\NgocLinh\snapgame\ivygames\GameCasual_ManHinhDOc.xlsx')
ws = wb.active

# Column mapping (1-indexed):
# Col 1: STT
# Col 2: Name (Tên)
# Col 3: Short Name
# Col 4: Mã code
# Col 5: Slug
# Col 6: Link
# Col 7: Mô tả (shortDescription)
# Col 8: Hướng dẫn chơi (description)

list_games = []

# Data starts from row 5, header rows are 1-4
for row in range(5, ws.max_row + 1):
    stt = ws.cell(row=row, column=1).value
    if stt is None:
        break
    
    name = ws.cell(row=row, column=2).value
    short_name = ws.cell(row=row, column=3).value
    code = ws.cell(row=row, column=4).value
    slug = ws.cell(row=row, column=5).value
    link = ws.cell(row=row, column=6).value
    short_description = ws.cell(row=row, column=7).value
    description_raw = ws.cell(row=row, column=8).value
    
    # Clean up name (remove trailing spaces)
    if name:
        name = name.strip()
    if short_name:
        short_name = short_name.strip()
    if code:
        code = code.strip()
    if slug:
        slug = slug.strip()
    if link:
        link = link.strip()
    if short_description:
        short_description = short_description.strip()
    
    # Format description as HTML like the example
    if description_raw:
        description_raw = description_raw.strip()
        # Split by newlines and wrap in <p> tags
        lines = description_raw.split('\n')
        html_parts = []
        for line in lines:
            line = line.strip()
            if line:
                # Replace double spaces with &nbsp;
                line = line.replace('  ', '&nbsp;')
                html_parts.append(f'<p>{line}</p>')
        description_html = ''.join(html_parts)
    else:
        description_html = ''
    
    game = {
        "price": 0,
        "type": "G4",
        "status": "ACTIVE",
        "developer": 1,
        "name": name,
        "slug": slug,
        "shortName": short_name,
        "categoryId": 12,
        "code": code,
        "link": link,
        "shortDescription": short_description,
        "description": description_html,
        "categoryIds": [12],
        "servers": [],
        "builds": [],
        "versions": [],
        "events": []
    }
    
    list_games.append(game)

# Wrap in the list_games structure
output = {
    "list_games": list_games
}

# Write JSON file
output_path = r'D:\NgocLinh\snapgame\ivygames\list_games.json'
with open(output_path, 'w', encoding='utf-8') as f:
    json.dump(output, f, ensure_ascii=False, indent=4)

print(f'Successfully generated {len(list_games)} game objects')
print(f'Output file: {output_path}')
