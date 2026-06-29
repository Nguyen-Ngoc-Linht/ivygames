# -*- coding: utf-8 -*-
"""
Script thêm 29 game bổ sung vào file GameCasual_ManHinhDOc.xlsx
"""
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from copy import copy
import os

# ========== CẤU HÌNH ==========
EXCEL_FILE = os.path.join(os.path.dirname(os.path.dirname(__file__)), "GameCasual_ManHinhDOc.xlsx")
DRIVE_LINK = "https://drive.google.com/drive/folders/1BMe639s1BhUuUgNDWe81aHbDnmlM33b5?usp=drive_link"
BASE_URL = "https://ivygames.pages.dev"

# Hướng dẫn chơi mặc định
DEFAULT_INSTRUCTION = (
    "Trò chơi vô cùng đơn giản:\n"
    "+ Điện thoại: Nhấn vào màn hình.\n"
    "+ Máy tính: sử dụng chuột trái  nhấp vào màn hình\n\n"
)

# ========== DANH SÁCH 29 GAME MỚI ==========
# Format: (game_name, short_name, code, slug, folder_build_name, description)
NEW_GAMES = [
    (
        "Jago",
        "Jago",
        "JGGAME",
        "jago",
        "JAGO",
        "Jago là một tựa game giải trí hấp dẫn với lối chơi nhanh và đầy thử thách. "
        "Người chơi sẽ tham gia vào những màn chơi đa dạng, đòi hỏi sự tập trung và phản xạ tốt "
        "để vượt qua các chướng ngại vật và đạt điểm số cao. Với đồ họa sinh động cùng cơ chế chơi cuốn hút, "
        "trò chơi phù hợp để thư giãn và giải trí mọi lúc mọi nơi."
    ),
    (
        "Air Warfare",
        "Air War",
        "AWGAME",
        "air-warfare",
        "Air-Warfare",
        "Air Warfare là một tựa game hành động không chiến hấp dẫn với phong cách chơi nhanh và kịch tính. "
        "Người chơi sẽ điều khiển máy bay chiến đấu, xuyên qua làn đạn và tiêu diệt kẻ thù trên bầu trời. "
        "Với đồ họa sống động và nhịp độ chơi cuốn hút, trò chơi mang đến trải nghiệm giải trí đầy phấn khích."
    ),
    (
        "Fruit Slasher",
        "Fruit Slash",
        "FSGAME",
        "fruit-slasher",
        "Fruit-Slasher",
        "Fruit Slasher là một tựa game giải trí với lối chơi chém hoa quả quen thuộc và đầy cuốn hút. "
        "Người chơi sẽ sử dụng ngón tay hoặc chuột để chém những trái cây bay lên màn hình, "
        "tránh chém phải bom và cố gắng đạt điểm số cao nhất. Với hình ảnh tươi sáng và nhịp chơi sôi động, "
        "trò chơi phù hợp cho mọi lứa tuổi."
    ),
    (
        "Stick Monkey",
        "Stick Monkey",
        "SMGAME",
        "stick-monkey",
        "Stick-Monkey",
        "Stick Monkey là một tựa game phiêu lưu vui nhộn với nhân vật khỉ đáng yêu. "
        "Người chơi sẽ điều khiển chú khỉ vượt qua các chướng ngại vật, nhảy qua các nền tảng "
        "và thu thập vật phẩm dọc đường. Với lối chơi đơn giản nhưng gây nghiện, "
        "trò chơi mang đến những phút giây thư giãn thú vị."
    ),
    (
        "Brick Out",
        "Brick Out",
        "BOGAME",
        "brick-out",
        "Brick-Out",
        "Brick Out là một tựa game phá gạch kinh điển với phong cách hiện đại đầy cuốn hút. "
        "Người chơi sẽ điều khiển thanh trượt để bắn bóng phá vỡ các viên gạch trên màn hình, "
        "vượt qua nhiều cấp độ với độ khó tăng dần. Với lối chơi quen thuộc nhưng đầy thử thách, "
        "trò chơi phù hợp cho mọi đối tượng người chơi."
    ),
    (
        "Fruit Snake",
        "Fruit Snake",
        "FNGAME",
        "fruit-snake",
        "Fruit-Snake",
        "Fruit Snake là một tựa game giải trí lấy cảm hứng từ trò chơi rắn săn mồi kinh điển. "
        "Người chơi sẽ điều khiển chú rắn di chuyển trên màn hình để ăn các loại trái cây, "
        "cơ thể rắn sẽ dài ra sau mỗi lần ăn. Với đồ họa tươi sáng và lối chơi gây nghiện, "
        "trò chơi mang đến trải nghiệm giải trí nhẹ nhàng và thú vị."
    ),
    (
        "Tank Defender",
        "Tank Defense",
        "TDGAME2",
        "tank-defender",
        "Tank-Defender",
        "Tank Defender là một tựa game hành động phòng thủ hấp dẫn. "
        "Người chơi sẽ điều khiển xe tăng chiến đấu, bảo vệ căn cứ trước những đợt tấn công của kẻ thù. "
        "Với nhiều loại vũ khí và cấp độ đa dạng, trò chơi đòi hỏi chiến thuật và phản xạ nhanh nhạy. "
        "Đồ họa sống động mang đến trải nghiệm chiến đấu đầy kịch tính."
    ),
    (
        "Fishing Frenzy",
        "Fish Frenzy",
        "FFGAME",
        "fishing-frenzy",
        "Fishing-Frenzy",
        "Fishing Frenzy là một tựa game giải trí với chủ đề câu cá vui nhộn và thư giãn. "
        "Người chơi sẽ thả câu và bắt những con cá đa dạng trong các vùng nước khác nhau, "
        "cố gắng thu thập được nhiều cá nhất có thể. Với hình ảnh sinh động và lối chơi nhẹ nhàng, "
        "trò chơi phù hợp để thư giãn trong mọi thời điểm."
    ),
    (
        "Space Purge",
        "Space Purge",
        "SPGAME",
        "space-purge",
        "Space-Purge",
        "Space Purge là một tựa game bắn súng không gian hấp dẫn với nhịp độ nhanh. "
        "Người chơi sẽ điều khiển phi thuyền chiến đấu, tiêu diệt các đợt kẻ thù ngoài hành tinh "
        "và vượt qua nhiều cấp độ đầy thử thách. Với đồ họa vũ trụ ấn tượng và hiệu ứng bắn phá mãn nhãn, "
        "trò chơi mang đến trải nghiệm hành động đầy phấn khích."
    ),
    (
        "Super Cowboy Run",
        "Cowboy Run",
        "SCRGAME",
        "super-cowboy-run",
        "Super-Cowboy-Run",
        "Super Cowboy Run là một tựa game chạy vô tận với chủ đề cao bồi miền Tây hoang dã. "
        "Người chơi sẽ điều khiển chàng cao bồi chạy, nhảy và né tránh các chướng ngại vật, "
        "thu thập vàng và vật phẩm dọc đường. Với phong cách đồ họa vui nhộn và lối chơi cuốn hút, "
        "trò chơi phù hợp cho mọi lứa tuổi."
    ),
    (
        "Duck Shooter",
        "Duck Shoot",
        "DSGAME",
        "duck-shooter",
        "Duck-Shooter",
        "Duck Shooter là một tựa game bắn súng giải trí với lối chơi đơn giản nhưng cuốn hút. "
        "Người chơi sẽ ngắm bắn những chú vịt bay qua màn hình, cố gắng đạt điểm số cao nhất "
        "với số đạn giới hạn. Với phong cách đồ họa hoạt hình vui nhộn và cơ chế chơi dễ tiếp cận, "
        "trò chơi mang đến những phút giây thư giãn nhẹ nhàng."
    ),
    (
        "Christmas Balloons",
        "Xmas Balloon",
        "CBGAME",
        "christmas-balloons",
        "Christmas-Balloons",
        "Christmas Balloons là một tựa game giải trí mang không khí Giáng Sinh vui nhộn và ấm áp. "
        "Người chơi sẽ tham gia vào những thử thách bắn bóng bay với chủ đề lễ hội, "
        "thu thập quà tặng và đạt điểm số cao. Với đồ họa đầy màu sắc Noel và nhạc nền rộn ràng, "
        "trò chơi mang đến trải nghiệm giải trí tuyệt vời trong mùa lễ hội."
    ),
    (
        "Great Air Battles",
        "Air Battles",
        "GABGAME",
        "great-air-battles",
        "Great-Air-Battles",
        "Great Air Battles là một tựa game hành động không chiến hoành tráng. "
        "Người chơi sẽ tham gia vào những trận đánh trên không kịch tính, điều khiển máy bay "
        "né tránh đạn và tiêu diệt kẻ thù. Với nhiều loại máy bay và cấp độ đa dạng, "
        "trò chơi mang đến trải nghiệm chiến đấu đầy phấn khích và thử thách."
    ),
    (
        "Professor Bubble",
        "Prof Bubble",
        "PBGAME2",
        "professor-bubble",
        "Professor-Bubble",
        "Professor Bubble là một tựa game bắn bong bóng đầy trí tuệ và cuốn hút. "
        "Người chơi sẽ bắn những quả bong bóng màu sắc để ghép nhóm và phá vỡ chúng, "
        "vượt qua nhiều cấp độ với độ khó tăng dần. Với lối chơi chiến thuật nhẹ nhàng "
        "và đồ họa sinh động, trò chơi phù hợp cho mọi lứa tuổi."
    ),
    (
        "Ninja Run",
        "Ninja Run",
        "NRGAME2",
        "ninja-run",
        "Ninja-Run",
        "Ninja Run là một tựa game chạy vô tận với chủ đề ninja đầy kịch tính. "
        "Người chơi sẽ điều khiển chiến binh ninja chạy, nhảy và né tránh các bẫy nguy hiểm, "
        "thu thập vật phẩm và tiêu diệt kẻ thù trên đường đi. Với lối chơi nhanh và phản xạ cao, "
        "trò chơi mang đến trải nghiệm hành động đầy hấp dẫn."
    ),
    (
        "Balloon Paradise",
        "Balloon Fun",
        "BPGAME2",
        "balloon-paradise",
        "Balloon-Paradise",
        "Balloon Paradise là một tựa game giải trí với thế giới bóng bay đầy màu sắc và vui nhộn. "
        "Người chơi sẽ tham gia vào những thử thách hấp dẫn, bắn và ghép các bóng bay "
        "để vượt qua nhiều cấp độ thú vị. Với đồ họa tươi sáng và lối chơi nhẹ nhàng, "
        "trò chơi phù hợp để thư giãn và giải trí."
    ),
    (
        "Hot Jewels",
        "Hot Jewels",
        "HJGAME",
        "hot-jewels",
        "Hot-Jewels",
        "Hot Jewels là một tựa game xếp kim cương hấp dẫn với lối chơi match-3 kinh điển. "
        "Người chơi sẽ hoán đổi vị trí các viên đá quý lấp lánh để tạo thành hàng 3 hoặc nhiều hơn, "
        "thu thập điểm số và vượt qua các cấp độ đầy thử thách. Với hiệu ứng đẹp mắt "
        "và lối chơi gây nghiện, trò chơi phù hợp cho mọi đối tượng."
    ),
    (
        "Smiles",
        "Smiles",
        "SLGAME",
        "smiles",
        "Smiles",
        "Smiles là một tựa game giải trí vui nhộn với phong cách đồ họa đáng yêu. "
        "Người chơi sẽ tham gia vào những thử thách ghép nối các biểu tượng mặt cười, "
        "vượt qua nhiều cấp độ với độ khó tăng dần. Với lối chơi đơn giản nhưng cuốn hút "
        "và hình ảnh tươi sáng, trò chơi mang đến những phút giây thư giãn đầy niềm vui."
    ),
    (
        "Zombie Uprising",
        "Zombie Rise",
        "ZUGAME",
        "zombie-uprising",
        "Zombie-Uprising",
        "Zombie Uprising là một tựa game hành động sinh tồn với chủ đề zombie đầy kịch tính. "
        "Người chơi sẽ chiến đấu chống lại các đợt tấn công của zombie, sử dụng vũ khí "
        "và chiến thuật để bảo vệ bản thân. Với nhiều loại zombie và cấp độ đa dạng, "
        "trò chơi mang đến trải nghiệm hành động đầy căng thẳng và hấp dẫn."
    ),
    (
        "Flappy Ball",
        "Flappy Ball",
        "FBGAME3",
        "flappy-ball",
        "Flappy-Ball",
        "Flappy Ball là một tựa game giải trí với lối chơi flappy quen thuộc nhưng đầy thử thách. "
        "Người chơi sẽ điều khiển quả bóng bay qua các chướng ngại vật bằng cách nhấn màn hình, "
        "cố gắng đi được xa nhất có thể. Với lối chơi đơn giản nhưng gây nghiện, "
        "trò chơi phù hợp để thử thách bản thân mọi lúc mọi nơi."
    ),
    (
        "Flappy Bounce",
        "Flappy Bounce",
        "FBNGAME",
        "flappy-bounce",
        "Flappy-Bounce",
        "Flappy Bounce là một tựa game giải trí kết hợp lối chơi flappy và nảy bóng đầy cuốn hút. "
        "Người chơi sẽ điều khiển vật thể nảy qua các chướng ngại vật, đòi hỏi sự tập trung "
        "và phản xạ nhanh. Với đồ họa sinh động và lối chơi đơn giản nhưng thử thách, "
        "trò chơi mang đến trải nghiệm giải trí nhẹ nhàng và thú vị."
    ),
    (
        "Bashorun",
        "Bashorun",
        "BRGAME2",
        "bashorun",
        "Bashorun",
        "Bashorun là một tựa game chạy vô tận với phong cách hành động và phiêu lưu hấp dẫn. "
        "Người chơi sẽ điều khiển nhân vật chạy, nhảy và né tránh các chướng ngại vật, "
        "thu thập vật phẩm và đạt điểm số cao nhất. Với lối chơi nhanh và đồ họa sinh động, "
        "trò chơi mang đến trải nghiệm giải trí đầy kịch tính."
    ),
    (
        "Splishy Fish",
        "Splish Fish",
        "SFGAME",
        "splishy-fish",
        "Splishy-Fish",
        "Splishy Fish là một tựa game giải trí với chủ đề cá vui nhộn và dễ thương. "
        "Người chơi sẽ điều khiển chú cá bơi qua các chướng ngại vật dưới nước, "
        "thu thập vật phẩm và cố gắng đi được xa nhất có thể. Với đồ họa đáng yêu "
        "và lối chơi cuốn hút, trò chơi phù hợp cho mọi lứa tuổi."
    ),
    (
        "Scary Run",
        "Scary Run",
        "SCGAME",
        "scary-run",
        "Scary-Run",
        "Scary Run là một tựa game chạy vô tận với chủ đề kinh dị đầy kịch tính. "
        "Người chơi sẽ điều khiển nhân vật chạy trốn trong bóng tối, né tránh các quái vật "
        "và bẫy nguy hiểm trên đường đi. Với bầu không khí rùng rợn và lối chơi nhanh, "
        "trò chơi mang đến trải nghiệm hành động đầy hồi hộp."
    ),
    (
        "Monsters Match 3",
        "Monster Match",
        "MMGAME",
        "monsters-match-3",
        "Monsters-Match-3",
        "Monsters Match 3 là một tựa game xếp hình match-3 với chủ đề quái vật vui nhộn. "
        "Người chơi sẽ hoán đổi vị trí các quái vật đáng yêu để tạo thành hàng 3 hoặc nhiều hơn, "
        "thu thập điểm số và vượt qua nhiều cấp độ đa dạng. Với đồ họa hoạt hình sinh động "
        "và lối chơi gây nghiện, trò chơi phù hợp cho mọi đối tượng."
    ),
    (
        "Math Game For Kids",
        "Math Kids",
        "MKGAME",
        "math-game-for-kids",
        "Math-Game-For-Kids",
        "Math Game For Kids là một tựa game giáo dục giúp trẻ em học toán một cách vui nhộn. "
        "Người chơi sẽ giải các bài toán cộng, trừ, nhân, chia với độ khó tăng dần, "
        "thu thập phần thưởng và đạt thành tích cao. Với giao diện thân thiện và hình ảnh đáng yêu, "
        "trò chơi giúp trẻ phát triển tư duy toán học một cách tự nhiên."
    ),
    (
        "Donut Crash Saga",
        "Donut Crash",
        "DCGAME",
        "donut-crash-saga",
        "Donut-Crash-Saga",
        "Donut Crash Saga là một tựa game match-3 với chủ đề bánh donut ngọt ngào và hấp dẫn. "
        "Người chơi sẽ ghép nối các chiếc bánh donut nhiều màu sắc để phá vỡ chúng, "
        "vượt qua nhiều cấp độ với mục tiêu đa dạng. Với đồ họa tươi sáng và lối chơi gây nghiện, "
        "trò chơi mang đến trải nghiệm giải trí ngọt ngào và thú vị."
    ),
    (
        "Stick Panda",
        "Stick Panda",
        "SPGAME2",
        "stick-panda",
        "Stick-Panda",
        "Stick Panda là một tựa game phiêu lưu với nhân vật gấu trúc đáng yêu. "
        "Người chơi sẽ điều khiển chú gấu trúc vượt qua các chướng ngại vật, "
        "nhảy qua các nền tảng và thu thập vật phẩm trên đường đi. Với lối chơi đơn giản "
        "và đồ họa hoạt hình dễ thương, trò chơi phù hợp cho mọi lứa tuổi."
    ),
    (
        "Crazy Car",
        "Crazy Car",
        "CCGAME2",
        "crazy-car",
        "Crazy-Car",
        "Crazy Car là một tựa game đua xe với lối chơi nhanh và đầy kịch tính. "
        "Người chơi sẽ điều khiển chiếc xe lao nhanh trên đường, né tránh các phương tiện khác "
        "và thu thập vật phẩm để đạt điểm số cao. Với đồ họa sống động và nhịp độ chơi cuốn hút, "
        "trò chơi mang đến trải nghiệm đua xe đầy phấn khích."
    ),
]


def add_games():
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb.active

    # Tìm dòng cuối cùng có dữ liệu
    last_data_row = 4
    for r in range(5, ws.max_row + 1):
        if ws.cell(row=r, column=1).value is not None:
            last_data_row = r

    last_stt = ws.cell(row=last_data_row, column=1).value
    print(f"Dòng cuối cùng có dữ liệu: Row {last_data_row}, STT = {last_stt}")

    # Lấy style từ dòng dữ liệu cuối cùng làm mẫu
    template_row = last_data_row
    template_styles = {}
    for col in range(1, 11):
        cell = ws.cell(row=template_row, column=col)
        template_styles[col] = {
            'font': copy(cell.font),
            'alignment': copy(cell.alignment),
            'border': copy(cell.border),
            'fill': copy(cell.fill),
            'number_format': cell.number_format,
        }

    # Thêm 29 game mới
    start_row = last_data_row + 1
    for i, (name, short_name, code, slug, folder_build, desc) in enumerate(NEW_GAMES):
        row = start_row + i
        stt = last_stt + 1 + i

        # Link deploy
        link = f"{BASE_URL}/{folder_build}/"

        # Data cho mỗi cột
        row_data = {
            1: stt,                    # STT
            2: name,                   # Tên (Name)
            3: short_name,             # Short Name
            4: code,                   # Mã code
            5: slug,                   # Slug
            6: link,                   # Link
            7: desc,                   # Mô tả
            8: DEFAULT_INSTRUCTION,    # Hướng dẫn chơi
            9: DRIVE_LINK,             # Link lưu file
            10: folder_build,          # Logo (1:1) - tên folder build
        }

        for col in range(1, 11):
            cell = ws.cell(row=row, column=col)
            cell.value = row_data[col]
            # Áp dụng style từ template
            style = template_styles[col]
            cell.font = copy(style['font'])
            cell.alignment = copy(style['alignment'])
            cell.border = copy(style['border'])
            cell.fill = copy(style['fill'])
            cell.number_format = style['number_format']

        print(f"  Đã thêm Row {row}: STT={stt}, Name='{name}', Code='{code}', Slug='{slug}'")

    # Lưu file
    wb.save(EXCEL_FILE)
    print(f"\n✅ Đã thêm {len(NEW_GAMES)} game mới vào file Excel!")
    print(f"   STT từ {last_stt + 1} đến {last_stt + len(NEW_GAMES)}")
    print(f"   Rows từ {start_row} đến {start_row + len(NEW_GAMES) - 1}")


if __name__ == "__main__":
    add_games()
