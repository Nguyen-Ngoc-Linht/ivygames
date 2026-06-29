import openpyxl
import json

wb = openpyxl.load_workbook(r'D:\NgocLinh\snapgame\ivygames\GameCasual_ManHinhDOc.xlsx')
ws = wb.active

output_lines = []
output_lines.append(f'Sheet name: {ws.title}')
output_lines.append(f'Max row: {ws.max_row}')
output_lines.append(f'Max col: {ws.max_column}')
output_lines.append('')

# Print header rows to find actual headers
for row in range(1, min(10, ws.max_row + 1)):
    vals = []
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row=row, column=col).value
        vals.append(str(v) if v is not None else '')
    output_lines.append(f'Row {row}: ' + ' | '.join(vals))

output_lines.append('')
output_lines.append('--- ALL DATA ROWS ---')

# Print all rows
for row in range(1, ws.max_row + 1):
    vals = []
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row=row, column=col).value
        vals.append(str(v) if v is not None else '')
    # Skip completely empty rows
    if any(v.strip() for v in vals):
        output_lines.append(f'Row {row}: ' + ' | '.join(vals))

with open(r'D:\NgocLinh\snapgame\ivygames\tools\excel_output.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(output_lines))

print('Done! Output written to excel_output.txt')
